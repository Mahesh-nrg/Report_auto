"""Blocks to a DOCX with python-docx: real Title / Heading n / List / Caption / Code styles,
tables with merged cells, and images cropped from the rendered page by ``pdftoppm``."""

from __future__ import annotations

import io
import shutil
from collections.abc import Sequence
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from . import limits
from .blocks import Block
from .extract import XML_BAD, Box, ToolError, run_tool
from .tables import Table

MARGIN_PT = 54.0
CODE_STYLE = "Code"
CODE_FONT = "Liberation Mono"
CODE_SIZE = 9
LIST_STYLES = {False: "List Bullet", True: "List Number"}


@dataclass(frozen=True)
class EmitCounts:
    tables: int = 0
    images: int = 0
    lists: int = 0
    images_capped: bool = False


def clean(text: str) -> str:
    return XML_BAD.sub("", text)


def crop(pdf: Path, page: int, box: Box, target: Path) -> Path | None:
    """PNG of one placed region (images plus vector overlays) or None."""
    scale = limits.CROP_DPI / 72
    x, y = int(box[0] * scale), int(box[1] * scale)
    width, height = int((box[2] - box[0]) * scale), int((box[3] - box[1]) * scale)
    if width < 4 or height < 4 or shutil.which("pdftoppm") is None:
        return None
    argv = [
        "pdftoppm", "-q", "-png", "-r", str(limits.CROP_DPI), "-f", str(page), "-l", str(page),
        "-x", str(x), "-y", str(y), "-W", str(width), "-H", str(height), "-singlefile",
        str(pdf), str(target),
    ]  # fmt: skip
    try:
        run_tool(argv)
    except ToolError:
        return None
    out = target.with_suffix(".png")
    return out if out.is_file() else None


def crop_all(pdf: Path, blocks: Sequence[Block], work: Path) -> dict[int, Path]:
    """Index of the image block -> PNG, for at most ``MAX_IMAGES`` blocks (pool of 4)."""
    wanted = [(i, b) for i, b in enumerate(blocks) if b.kind == "image" and b.bbox]
    wanted = wanted[: limits.MAX_IMAGES]
    work.mkdir(parents=True, exist_ok=True)
    with ThreadPoolExecutor(max_workers=limits.CROP_WORKERS) as pool:
        jobs = {
            i: pool.submit(crop, pdf, b.page, b.bbox, work / f"img{i:04d}")
            for i, b in wanted
            if b.bbox is not None
        }
        return {i: path for i, job in jobs.items() if (path := job.result()) is not None}


def _code_style(document: Any) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt

    names = {style.name for style in document.styles}
    if CODE_STYLE not in names:
        style = document.styles.add_style(CODE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.name = CODE_FONT
        style.font.size = Pt(CODE_SIZE)


def _page_setup(document: Any, size: tuple[float, float]) -> float:
    from docx.shared import Pt

    section = document.sections[0]
    section.page_width, section.page_height = Pt(size[0]), Pt(size[1])
    for name in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, name, Pt(MARGIN_PT))
    return max(72.0, size[0] - 2 * MARGIN_PT)


def _table(document: Any, table: Table) -> None:
    grid = document.add_table(rows=table.rows, cols=table.cols)
    grid.style = "Table Grid"
    for r in range(table.rows):
        for c in range(table.cols):
            rows, cols = table.spans[r][c]
            if rows == 0:
                continue
            cell = grid.cell(r, c)
            if rows > 1 or cols > 1:
                last_r, last_c = min(table.rows, r + rows) - 1, min(table.cols, c + cols) - 1
                cell = cell.merge(grid.cell(last_r, last_c))
            cell.text = clean(table.cells[r][c])


def _code(document: Any, text: str) -> None:
    paragraph = document.add_paragraph(style=CODE_STYLE)
    for index, part in enumerate(text.split("\n")):
        run = paragraph.add_run()
        if index:
            run.add_break()
        run.text = clean(part)


def _picture(document: Any, block: Block, path: Path | None, text_width: float) -> bool:
    from docx.shared import Pt

    if path is None or block.bbox is None:
        return False
    width = min(block.bbox[2] - block.bbox[0], text_width)
    try:
        document.add_picture(str(path), width=Pt(width))
    except Exception:  # noqa: BLE001 - a broken crop is skipped, not fatal
        return False
    return True


def _write_block(document: Any, block: Block) -> None:
    text = clean(block.text)
    if block.kind == "title":
        document.add_heading(text, level=0)
    elif block.kind == "heading":
        document.add_heading(text, level=max(1, min(block.level, 4)))
    elif block.kind == "list":
        document.add_paragraph(text, style=LIST_STYLES[block.ordered])
    elif block.kind == "code":
        _code(document, block.text)
    elif block.kind == "caption":
        document.add_paragraph(text, style="Caption")
    elif text:
        document.add_paragraph(text)


def write_docx(
    blocks: Sequence[Block], pdf: Path, work: Path, page_size: tuple[float, float]
) -> tuple[bytes, EmitCounts]:
    import docx

    document = docx.Document()
    _code_style(document)
    text_width = _page_setup(document, page_size)
    crops = crop_all(pdf, blocks, work / "crops")
    tables = images = lists = 0
    for index, block in enumerate(blocks):
        if block.kind == "table" and block.table is not None:
            _table(document, block.table)
            tables += 1
        elif block.kind == "image":
            images += _picture(document, block, crops.get(index), text_width)
        else:
            lists += block.kind == "list"
            _write_block(document, block)
    capped = sum(1 for b in blocks if b.kind == "image") > limits.MAX_IMAGES
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), EmitCounts(tables, images, lists, capped)


__all__ = ["EmitCounts", "clean", "crop", "write_docx"]
